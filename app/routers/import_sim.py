import io, re, os, calendar, json, uuid
from datetime import date, datetime
from typing import Optional

import openpyxl
import xlrd

from fastapi import APIRouter, Request, Depends, UploadFile, File, Form
from fastapi.responses import HTMLResponse, RedirectResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.simcard import SimCard, ActivityLog

router   = APIRouter(prefix="/import", tags=["import"])
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
templates = Jinja2Templates(directory=os.path.join(BASE_DIR, "templates"))

# Temp folder for preview data (avoids session cookie size limit)
TMP_DIR = os.path.join(BASE_DIR, "..", "tmp_import")
os.makedirs(TMP_DIR, exist_ok=True)


def _save_preview(data: list) -> str:
    """Save preview rows to a temp JSON file. Returns token."""
    token = str(uuid.uuid4())
    path  = os.path.join(TMP_DIR, f"{token}.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, default=str)
    return token


def _load_preview(token: str) -> list:
    """Load preview rows from temp file."""
    if not token:
        return []
    path = os.path.join(TMP_DIR, f"{token}.json")
    if not os.path.exists(path):
        return []
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def _delete_preview(token: str):
    path = os.path.join(TMP_DIR, f"{token}.json")
    try:
        os.remove(path)
    except Exception:
        pass


# ── Helpers ───────────────────────────────────────────────────

def _normalize_carrier(raw) -> str:
    if raw is None:
        return "other"
    s = str(raw).strip().upper()
    if "AIS" in s:
        return "AIS"
    if "TRUE" in s or "DTAC" in s:
        return "TRUE"
    if s == "NT":
        return "NT"
    return str(raw).strip() or "other"


def _extract_phone(raw) -> str:
    if raw is None:
        return ""
    if isinstance(raw, (int, float)):
        s = str(int(raw))
        return ("0" + s) if len(s) == 9 else s
    s = str(raw).strip()
    m = re.search(r'0[6-9]\d{8}', s)
    if m:
        return m.group()
    digits = re.sub(r'[^0-9]', '', s)
    if len(digits) >= 9:
        return ("0" + digits)[-10:] if not digits.startswith("0") else digits[-10:]
    return s


def _extract_iccid(raw) -> str:
    if raw is None:
        return ""
    if isinstance(raw, float):
        raw = int(raw)
    return str(raw).strip().replace(".0", "")


def _safe_date(yr: int, mo: int, day: int) -> Optional[date]:
    if yr > 2400:
        yr -= 543
    try:
        return date(yr, mo, day)
    except ValueError:
        try:
            last = calendar.monthrange(yr, mo)[1]
            return date(yr, mo, last)
        except Exception:
            return None


def _parse_date_str(raw) -> Optional[date]:
    if raw is None:
        return None
    if isinstance(raw, datetime):
        d = raw.date()
        return date(d.year - 543, d.month, d.day) if d.year > 2400 else d
    if isinstance(raw, date):
        return date(raw.year - 543, raw.month, raw.day) if raw.year > 2400 else raw
    if isinstance(raw, (int, float)):
        return None
    s = re.sub(r'^(START|EXP|Start|start|exp)\s*', '', str(raw).strip(), flags=re.IGNORECASE).strip()
    if not s:
        return None
    m = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{4})$', s)
    if m:
        return _safe_date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    m = re.match(r'^(\d{1,2})/(\d{4})$', s)
    if m:
        return _safe_date(int(m.group(2)), int(m.group(1)), 1)
    m = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{2})$', s)
    if m:
        return _safe_date(2500 + int(m.group(3)), int(m.group(2)), int(m.group(1)))
    m = re.match(r'^(\d{1,2})/(\d{2})$', s)
    if m:
        return _safe_date(2500 + int(m.group(2)), int(m.group(1)), 1)
    return None


def _build_plan(type_val, package_val) -> str:
    parts = [str(v).strip() for v in (type_val, package_val) if v]
    return " | ".join(parts)


def _build_notes(hardware, remark, project) -> str:
    parts = []
    if project:
        parts.append(f"[{project}]")
    if hardware:
        parts.append(f"HW:{hardware}")
    if remark:
        parts.append(str(remark).strip())
    return " | ".join(parts)


def parse_excel(file_bytes: bytes, filename: str) -> list:
    rows = []
    if filename.lower().endswith(".xls"):
        wb = xlrd.open_workbook(file_contents=file_bytes)
        ws = wb.sheet_by_index(0)
        headers = [str(ws.cell_value(0, c)).strip().lower() for c in range(ws.ncols)]
        for r in range(1, ws.nrows):
            vals = [ws.cell_value(r, c) for c in range(ws.ncols)]
            rows.append(dict(zip(headers, vals)))
    else:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active
        headers = [str(ws.cell(1, c).value or "").strip().lower()
                   for c in range(1, ws.max_column + 1)]
        for r in range(2, ws.max_row + 1):
            vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            if all(v is None for v in vals):
                continue
            rows.append(dict(zip(headers, vals)))
    return rows


def row_to_simcard(row: dict, row_num: int) -> dict:
    iccid      = _extract_iccid(row.get("s/n")) or f"SN-IMPORT-{row_num:04d}"
    phone      = _extract_phone(row.get("number"))
    carrier    = _normalize_carrier(row.get("network"))
    plan       = _build_plan(row.get("type"), row.get("package"))
    notes_txt  = _build_notes(row.get("hardware"), row.get("remark"), row.get("project"))
    activation = _parse_date_str(row.get("start"))
    expiry     = _parse_date_str(row.get("exp"))
    department = str(row.get("station") or "").strip()
    assigned   = str(row.get("postion") or "").strip()
    return {
        "iccid":           iccid,
        "phone_number":    phone or f"N/A-{row_num}",
        "carrier":         carrier,
        "status":          "active",
        "plan":            plan,
        "assigned_to":     assigned,
        "department":      department,
        "activation_date": str(activation) if activation else "",
        "expiry_date":     str(expiry)     if expiry     else "",
        "monthly_cost":    0,
        "notes":           notes_txt,
        "_row":            row_num,
        "_raw_number":     str(row.get("number") or ""),
        "_raw_start":      str(row.get("start")  or ""),
        "_raw_exp":        str(row.get("exp")    or ""),
        "_status":         "new",
    }


# ── Routes ────────────────────────────────────────────────────

@router.get("", response_class=HTMLResponse)
async def import_page(request: Request):
    if not request.session.get("user_id"):
        return RedirectResponse(url="/login", status_code=302)
    return templates.TemplateResponse("import.html", {
        "request": request, "preview": None, "stats": None, "error": None,
    })


@router.post("/preview", response_class=HTMLResponse)
async def preview_import(
    request: Request,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    if not request.session.get("user_id"):
        return RedirectResponse(url="/login", status_code=302)

    filename = file.filename or ""
    if not filename.lower().endswith((".xlsx", ".xls")):
        return templates.TemplateResponse("import.html", {
            "request": request, "preview": None, "stats": None,
            "error": "กรุณาอัปโหลดไฟล์ .xlsx หรือ .xls เท่านั้น",
        })
    try:
        content  = await file.read()
        raw_rows = parse_excel(content, filename)
        parsed   = [row_to_simcard(r, i + 1) for i, r in enumerate(raw_rows)
                    if any(v for v in r.values())]

        existing_iccids = {s.iccid        for s in db.query(SimCard.iccid).all()}
        existing_phones = {s.phone_number  for s in db.query(SimCard.phone_number).all()}

        stats = {"new": 0, "duplicate_iccid": 0, "duplicate_phone": 0, "skip": 0, "intra_dupe": 0}
        seen_iccids_in_file = set()
        for p in parsed:
            if not p["phone_number"] or p["phone_number"].startswith("N/A"):
                p["_status"] = "skip";            stats["skip"] += 1
            elif p["iccid"] in seen_iccids_in_file:
                p["_status"] = "skip";            stats["intra_dupe"] += 1
                stats["skip"] += 1
            elif p["iccid"] in existing_iccids:
                p["_status"] = "duplicate_iccid"; stats["duplicate_iccid"] += 1
            elif p["phone_number"] in existing_phones:
                p["_status"] = "duplicate_phone"; stats["duplicate_phone"] += 1
            else:
                p["_status"] = "new";             stats["new"] += 1
            seen_iccids_in_file.add(p["iccid"])

        # Save to temp file — avoid session cookie size limit
        token = _save_preview(parsed)
        request.session["import_token"]    = token
        request.session["import_filename"] = filename

        return templates.TemplateResponse("import.html", {
            "request":  request,
            "preview":  parsed,
            "stats":    stats,
            "filename": filename,
            "token":    token,
            "error":    None,
        })

    except Exception as e:
        import traceback
        return templates.TemplateResponse("import.html", {
            "request": request, "preview": None, "stats": None,
            "error": f"เกิดข้อผิดพลาด: {str(e)}\n{traceback.format_exc()}",
        })


@router.post("/confirm")
async def confirm_import(
    request: Request,
    mode:  str = Form("new_only"),
    token: str = Form(""),
    db: Session = Depends(get_db),
):
    if not request.session.get("user_id"):
        return RedirectResponse(url="/login", status_code=302)

    preview_data = _load_preview(token or request.session.get("import_token", ""))
    if not preview_data:
        return templates.TemplateResponse("import.html", {
            "request": request, "preview": None, "stats": None,
            "error": "ไม่พบข้อมูล Preview — กรุณาอัปโหลดไฟล์ใหม่อีกครั้ง",
        })

    def parse_d(val):
        try:
            return date.fromisoformat(val) if val else None
        except Exception:
            return None

    imported = skipped = updated = errors = 0
    # Track ICCIDs inserted in this batch to handle intra-file duplicates
    batch_iccids = set()

    for row in preview_data:
        try:
            status = row.get("_status", "skip")
            if status == "skip":
                skipped += 1
                continue

            iccid    = row["iccid"]
            act_date = parse_d(row.get("activation_date"))
            exp_date = parse_d(row.get("expiry_date"))

            # Skip intra-file ICCID duplicate
            if iccid in batch_iccids:
                skipped += 1
                continue

            if status == "duplicate_iccid" and mode == "overwrite_all":
                sim = db.query(SimCard).filter(SimCard.iccid == iccid).first()
                if sim:
                    sim.phone_number    = row["phone_number"]
                    sim.carrier         = row["carrier"]
                    sim.plan            = row["plan"]
                    sim.assigned_to     = row["assigned_to"]
                    sim.department      = row["department"]
                    sim.activation_date = act_date
                    sim.expiry_date     = exp_date
                    sim.notes           = row["notes"]
                    sim.updated_at      = datetime.utcnow()
                    db.flush()
                    updated += 1
                    batch_iccids.add(iccid)
                continue

            if status == "new":
                try:
                    db.add(SimCard(
                        iccid=iccid,
                        phone_number=row["phone_number"],
                        carrier=row["carrier"],
                        status="active",
                        plan=row["plan"],
                        assigned_to=row["assigned_to"],
                        department=row["department"],
                        activation_date=act_date,
                        expiry_date=exp_date,
                        monthly_cost=0,
                        notes=row["notes"],
                    ))
                    db.flush()   # flush per row — catch UNIQUE error early
                    imported += 1
                    batch_iccids.add(iccid)
                except Exception:
                    db.rollback()
                    skipped += 1   # ICCID already in DB (race condition)

        except Exception as e:
            errors += 1
            print(f"Import error row={row.get('_row')}: {e}")

    try:
        db.add(ActivityLog(
            user_id=request.session.get("user_id"),
            action=f"Import Excel: +{imported} update={updated} skip={skipped} err={errors}",
        ))
        db.commit()
    except Exception as e:
        db.rollback()
        print(f"Final commit error: {e}")

    _delete_preview(token or request.session.get("import_token", ""))
    request.session.pop("import_token", None)

    request.session["import_result"] = {
        "imported": imported, "updated": updated,
        "skipped":  skipped,  "errors":  errors,
    }
    return RedirectResponse(url="/import/result", status_code=302)


@router.get("/result", response_class=HTMLResponse)
async def import_result(request: Request):
    if not request.session.get("user_id"):
        return RedirectResponse(url="/login", status_code=302)
    result = request.session.pop("import_result", None)
    return templates.TemplateResponse("import.html", {
        "request": request, "preview": None, "stats": None,
        "result":  result,  "error":  None,
    })

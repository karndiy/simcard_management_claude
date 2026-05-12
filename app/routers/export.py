import csv, io
from datetime import date
from fastapi import APIRouter, Request, Depends
from fastapi.responses import StreamingResponse, RedirectResponse
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from app.database import get_db
from app.models.simcard import SimCard

router = APIRouter(prefix="/export", tags=["export"])


def _get_sims(db, status=None, department=None):
    q = db.query(SimCard)
    if status:
        q = q.filter(SimCard.status == status)
    if department:
        q = q.filter(SimCard.department == department)
    return q.order_by(SimCard.carrier, SimCard.department).all()


# ── CSV Export ────────────────────────────────────────────────
@router.get("/csv")
async def export_csv(
    request: Request,
    status: str = "",
    department: str = "",
    db: Session = Depends(get_db),
):
    if not request.session.get("user_id"):
        return RedirectResponse(url="/login", status_code=302)

    sims = _get_sims(db, status or None, department or None)

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["ICCID", "เบอร์โทร", "ผู้ให้บริการ", "สถานะ", "แพ็กเกจ",
                     "ผู้ใช้งาน", "แผนก", "วันเปิดใช้", "วันหมดอายุ",
                     "ค่าบริการ/เดือน (บาท)", "หมายเหตุ"])
    for s in sims:
        writer.writerow([
            s.iccid, s.phone_number, s.carrier, s.status, s.plan or "",
            s.assigned_to or "", s.department or "",
            s.activation_date or "", s.expiry_date or "",
            float(s.monthly_cost or 0), s.notes or "",
        ])

    output.seek(0)
    filename = f"simcards_{date.today()}.csv"
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# ── Excel Export ──────────────────────────────────────────────
@router.get("/excel")
async def export_excel(
    request: Request,
    status: str = "",
    department: str = "",
    db: Session = Depends(get_db),
):
    if not request.session.get("user_id"):
        return RedirectResponse(url="/login", status_code=302)

    sims = _get_sims(db, status or None, department or None)

    wb = Workbook()
    ws = wb.active
    ws.title = "SIM Cards"

    # Header styling
    header_fill = PatternFill("solid", fgColor="1E40AF")
    header_font = Font(bold=True, color="FFFFFF")
    headers = ["ICCID", "เบอร์โทร", "ผู้ให้บริการ", "สถานะ", "แพ็กเกจ",
               "ผู้ใช้งาน", "แผนก", "วันเปิดใช้", "วันหมดอายุ",
               "ค่าบริการ/เดือน", "หมายเหตุ"]

    for col, h in enumerate(headers, 1):
        cell             = ws.cell(row=1, column=col, value=h)
        cell.font        = header_font
        cell.fill        = header_fill
        cell.alignment   = Alignment(horizontal="center")

    for row, s in enumerate(sims, 2):
        ws.append([
            s.iccid, s.phone_number, s.carrier, s.status, s.plan or "",
            s.assigned_to or "", s.department or "",
            str(s.activation_date or ""), str(s.expiry_date or ""),
            float(s.monthly_cost or 0), s.notes or "",
        ])

    # Auto column width
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"simcards_{date.today()}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
